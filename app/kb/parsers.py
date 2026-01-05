from __future__ import annotations

import csv
import io
import os
from PIL import Image

try:
    from docx import Document
except Exception:
    Document = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    import openpyxl
except Exception:
    openpyxl = None


def detect_ext(path: str) -> str:
    return (os.path.splitext(path)[1] or "").lower().strip(".")


def is_image_ext(ext: str) -> bool:
    return ext in {"png", "jpg", "jpeg", "webp"}


def parse_text_bytes(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


# Backward-compatible alias used by some modules
def parse_txt_bytes(data: bytes) -> str:
    return parse_text_bytes(data)


def parse_pdf_bytes(data: bytes) -> str:
    if PdfReader is None:
        return ""
    bio = io.BytesIO(data)
    reader = PdfReader(bio)
    parts = []
    for p in reader.pages:
        t = p.extract_text() or ""
        t = t.strip()
        if t:
            parts.append(t)
    return "\n\n".join(parts)


def parse_docx_bytes(data: bytes) -> str:
    if Document is None:
        return ""
    bio = io.BytesIO(data)
    doc = Document(bio)
    parts = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    return "\n".join(parts)


def parse_xlsx_bytes(data: bytes) -> str:
    if openpyxl is None:
        return ""
    bio = io.BytesIO(data)
    wb = openpyxl.load_workbook(bio, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"=== {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            cells = []
            for v in row:
                if v is None:
                    cells.append("")
                else:
                    cells.append(str(v))
            line = "\t".join(cells).strip()
            if line:
                parts.append(line)
    return "\n".join(parts).strip()


def parse_csv_bytes(data: bytes) -> str:
    text = parse_text_bytes(data)
    buf = io.StringIO(text)
    reader = csv.reader(buf)
    lines = []
    for row in reader:
        line = "\t".join([c.strip() for c in row]).strip()
        if line:
            lines.append(line)
    return "\n".join(lines).strip()


def parse_image_bytes_best_effort(data: bytes) -> str:
    # OCR не делаем. Просто пытаемся извлечь хоть какие-то метаданные/пиксельные признаки,
    # но по факту возвращаем пусто — файл будет пропущен индексатором.
    try:
        img = Image.open(io.BytesIO(data))
        _ = img.size  # noqa: F841
    except Exception:
        pass
    return ""
